"""
Biomedical Research Agent (Optimized + File Reading)
Embeds queries via BGE, retrieves hits from Qdrant,
synthesizes grounded research summaries via Qwen.
Supports direct file reading from /data/research-uploads/
"""

import os
import re
import csv
import logging
import httpx
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)

# ── Cluster & Timeout Config ──────────────────────────────────────────────────
EMBED_URL    = os.getenv("EMBED_URL",    "http://embedding-bge-svc.ai-models.svc.cluster.local:80/embed")
QDRANT_URL   = os.getenv("QDRANT_URL",   "http://qdrant-svc.ai-models.svc.cluster.local:6333")
UPLOAD_DIR   = Path(os.getenv("UPLOAD_DIR", "/data/research-uploads"))
HTTP_TIMEOUT = float(os.getenv("HTTP_TIMEOUT", "15.0"))
LLM_TIMEOUT  = float(os.getenv("LLM_TIMEOUT",  "120.0"))
QDRANT_TOP_K = int(os.getenv("QDRANT_TOP_K",   "5"))

# ── Context Budget ────────────────────────────────────────────────────────────
MAX_TOKENS        = int(os.getenv("MAX_TOKENS",       "512"))
CONTEXT_CHAR_CAP  = int(os.getenv("CONTEXT_CHAR_CAP", "1200"))
ABSTRACT_CHAR_CAP = int(os.getenv("ABSTRACT_CHAR_CAP", "200"))
FILE_CHAR_CAP     = int(os.getenv("FILE_CHAR_CAP",    "4000"))

SYSTEM_PROMPT = (
    "You are an expert biomedical Research Assistant with direct access to PubMed literature and BioModels. "
    "Synthesise evidence clearly using the retrieved context. "
    "Cite evidence directly using [PMID:XXXXX] or [Model:XXXXX]. "
    "Structure: Executive Summary → Key Evidence → Translational Implications → Limitations."
)

FILE_SYSTEM_PROMPT = (
    "You are an expert biomedical Research Assistant at ICCBS. "
    "Answer ONLY based on the provided file content. "
    "Be precise, specific, and extract exact values, durations, and outcomes mentioned. "
    "Do not add information not present in the file."
)

# ── File Path Detector ────────────────────────────────────────────────────────

FILE_PATH_REGEX = re.compile(
    r'(?:/data/research-uploads/)?([\w\-\. ]+\.(?:docx|txt|csv|tsv|md|pdf|fasta|fa|fna|vcf|gff|bed|xlsx))\b',
    re.IGNORECASE
)
def _extract_file_content(message: str) -> tuple[str, str]:
    """
    Detects file path reference in message and extracts content.
    Returns (filename, content) or ("", "") if no file referenced.
    """
    match = FILE_PATH_REGEX.search(message)
    if not match:
        return "", ""

    filename = match.group(1).strip()
    filepath = UPLOAD_DIR / filename

    if not filepath.exists():
        logger.warning("Referenced file not found: %s", filepath)
        return filename, f"[Error: File '{filename}' not found in /data/research-uploads/]"

    ext = filepath.suffix.lower()

    try:
        # ── Plain text ──
        if ext in {".txt", ".md"}:
            return filename, filepath.read_text(encoding="utf-8", errors="replace")[:FILE_CHAR_CAP]

        # ── Word documents ──
        elif ext == ".docx":
            try:
                from docx import Document
                doc = Document(filepath)
                content = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
                return filename, content[:FILE_CHAR_CAP]
            except ImportError:
                return filename, "[Error: python-docx not installed. Add 'python-docx' to requirements.txt]"

        # ── CSV / TSV ──
        elif ext in {".csv", ".tsv"}:
            delimiter = "\t" if ext == ".tsv" else ","
            with open(filepath, newline="", encoding="utf-8", errors="replace") as f:
                rows = list(csv.reader(f, delimiter=delimiter))
            formatted = "\n".join([delimiter.join(r) for r in rows[:100]])
            return filename, formatted[:FILE_CHAR_CAP]

        # ── FASTA / genomic ──
        elif ext in {".fasta", ".fa", ".fna", ".vcf", ".gff", ".bed"}:
            return filename, filepath.read_text(encoding="utf-8", errors="replace")[:FILE_CHAR_CAP]

	# ── PDF  ──
        elif ext == ".pdf":

            try:

                import pdfplumber

                with pdfplumber.open(filepath) as pdf:

                    pages_text = [page.extract_text() for page in pdf.pages[:10] if page.extract_text()]

                return filename, "\n\n".join(pages_text)[:FILE_CHAR_CAP]

            except ImportError:

                return filename, "[Error: pdfplumber not installed. Add 'pdfplumber' to requirements.txt]" 

        # ── Excel ──
        elif ext == ".xlsx":
            try:
                import openpyxl
                wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
                ws = wb.active
                rows = []
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i >= 100:
                        break
                    rows.append(",".join([str(c) if c is not None else "" for c in row]))
                return filename, "\n".join(rows)[:FILE_CHAR_CAP]
            except ImportError:
                return filename, "[Error: openpyxl not installed. Add 'openpyxl' to requirements.txt]"

        else:
            return filename, f"[Unsupported file type: {ext}]"

    except Exception as e:
        logger.error("File read error for %s: %s", filename, e)
        return filename, f"[Error reading '{filename}': {str(e)}]"


# ── Helper Services ──────────────────────────────────────────────────────────

async def _embed(text: str, client: httpx.AsyncClient) -> list[float] | None:
    try:
        resp = await client.post(
            EMBED_URL,
            json={"inputs": text[:512]},
            headers={"Content-Type": "application/json"},
            timeout=HTTP_TIMEOUT,
        )
        resp.raise_for_status()
        return resp.json()[0]
    except Exception as e:
        logger.error("BGE Embedding Service Failure: %s", e)
        return None


async def _qdrant_search(
    collection: str,
    vector: list[float],
    top_k: int,
    client: httpx.AsyncClient,
) -> list[dict[str, Any]]:
    try:
        resp = await client.post(
            f"{QDRANT_URL}/collections/{collection}/points/search",
            json={"vector": vector, "limit": top_k, "with_payload": True, "with_vector": False},
            timeout=HTTP_TIMEOUT,
        )
        resp.raise_for_status()
        return resp.json().get("result", [])
    except Exception as e:
        logger.error("Qdrant Vector Retrieval Error [%s]: %s", collection, e)
        return []


def _fmt_pubmed(hits: list[dict]) -> str:
    if not hits:
        return "No relevant PubMed articles found."
    lines = ["=== PUBMED RESEARCH EVIDENCE ==="]
    for i, hit in enumerate(hits, 1):
        p        = hit.get("payload", {})
        pmid     = p.get("pmid", "Unknown")
        title    = p.get("title", "Untitled")
        abstract = str(p.get("context", p.get("abstract", "No abstract available"))).strip()
        lines.append(f"[{i}] PMID:{pmid} - {title}\nAbstract: {abstract[:ABSTRACT_CHAR_CAP]}\n")
    return "\n".join(lines)


def _fmt_biomodels(hits: list[dict]) -> str:
    if not hits:
        return ""
    lines = ["=== BIOMODELS PATHWAY SYSTEMS ==="]
    for i, hit in enumerate(hits, 1):
        p        = hit.get("payload", {})
        model_id = p.get("model_id", "Unknown")
        desc     = str(p.get("context", p.get("description", "No description"))).strip()[:400]
        lines.append(f"[{i}] Model ID: {model_id}\nDescription: {desc}\n")
    return "\n".join(lines)


async def _call_llm(
    messages: list[dict],
    model_url: str,
    model_name: str,
    client: httpx.AsyncClient,
) -> str:
    resp = await client.post(
        f"{model_url}/v1/chat/completions",
        json={
            "model":      model_name,
            "messages":   messages,
            "max_tokens": MAX_TOKENS,
            "temperature": float(os.getenv("TEMPERATURE", "0.2")),
        },
        timeout=LLM_TIMEOUT,
    )
    resp.raise_for_status()
    return resp.json()["choices"][0]["message"]["content"]


# ── Execution Entrypoint ─────────────────────────────────────────────────────

async def run(
    message: str,
    history: list[dict],
    model_url: str,
    model_name: str,
    client: httpx.AsyncClient | None = None,
) -> dict[str, Any]:
    """
    RAG Orchestration Workflow:
    Step 0: File path detection → read file → answer directly (skips RAG)
    Step 1: Embed → Retrieve → Build Prompt → Generate (normal RAG path)
    """
    should_close = False
    if client is None:
        client = httpx.AsyncClient()
        should_close = True

    try:
        # ══════════════════════════════════════════════════════════════════
        # STEP 0: FILE READING PATH
        # ══════════════════════════════════════════════════════════════════
        filename, file_content = _extract_file_content(message)

        if filename and file_content:
            logger.info("File reference detected: %s (%d chars)", filename, len(file_content))

            if file_content.startswith("[Error"):
                return {
                    "answer": (
                        f"⚠️ **File Access Error**\n\n{file_content}\n\n"
                        "Please verify the file was uploaded via the Agent tab "
                        "and is stored in `/data/research-uploads/`."
                    ),
                    "sources":      [],
                    "agent":        "research",
                    "model":        model_name,
                    "context_hits": {"local_file": 0},
                }

            user_content = (
                f"File: `{filename}`\n"
                f"--- FILE CONTENT START ---\n"
                f"{file_content}\n"
                f"--- FILE CONTENT END ---\n\n"
                f"Question: {message}"
            )

            llm_messages = [
                {"role": "system", "content": FILE_SYSTEM_PROMPT},
                {"role": "user",   "content": user_content},
            ]

            try:
                answer = await _call_llm(llm_messages, model_url, model_name, client)
            except Exception as e:
                logger.error("LLM error during file-based answer: %s", e)
                answer = f"File read successfully but LLM inference failed: {e}"

            return {
                "answer":  answer,
                "sources": [{"type": "local_file", "filename": filename, "chars_read": len(file_content)}],
                "agent":   "research",
                "model":   model_name,
                "context_hits": {"local_file": 1},
            }

        # ══════════════════════════════════════════════════════════════════
        # STEP 1: NORMAL RAG PATH
        # ══════════════════════════════════════════════════════════════════

        # 1. Generate Vector Embeddings
        vector = await _embed(message, client)

        # 2. Vector Search
        pubmed_hits: list[dict]    = []
        biomodels_hits: list[dict] = []

        if vector:
            pubmed_hits    = await _qdrant_search("pubmed",    vector, QDRANT_TOP_K, client)
            biomodels_hits = await _qdrant_search("biomodels", vector, 3,            client)
            logger.info("Search Results - PubMed: %d, BioModels: %d", len(pubmed_hits), len(biomodels_hits))
        else:
            logger.warning("Retrieval bypassed: Embedding vector null.")

        # 3. Context Construction
        raw_context     = f"{_fmt_pubmed(pubmed_hits)}\n\n{_fmt_biomodels(biomodels_hits)}"
        bounded_context = raw_context.strip()[:CONTEXT_CHAR_CAP]

        grounded_msg = (
            f"Retrieved Research Context:\n{bounded_context}\n\n"
            f"User Question: {message[:1000]}\n\n"
            f"Provide a scientifically precise, cited synthesis based on the context above."
        )

        # 4. Message Assembly with History
        llm_messages = [{"role": "system", "content": SYSTEM_PROMPT}]
        for turn in history[-4:]:
            if turn.get("role") in ("user", "assistant"):
                llm_messages.append({"role": turn["role"], "content": turn["content"][:500]})
        llm_messages.append({"role": "user", "content": grounded_msg})

        # 5. LLM Synthesis
        try:
            answer = await _call_llm(llm_messages, model_url, model_name, client)
        except Exception as e:
            logger.error("vLLM Ingestion Error: %s", e)
            answer = (
                f"Data retrieved successfully ({len(pubmed_hits)} articles), "
                f"but LLM inference failed: {e}"
            )

        # 6. Structured Output
        sources = [
            {"type": "pubmed",    "pmid":     h.get("payload", {}).get("pmid"),     "title": h.get("payload", {}).get("title"), "score": round(h.get("score", 0), 3)}
            for h in pubmed_hits
        ] + [
            {"type": "biomodel",  "model_id": h.get("payload", {}).get("model_id"), "name":  h.get("payload", {}).get("name"),  "score": round(h.get("score", 0), 3)}
            for h in biomodels_hits
        ]

        return {
            "answer":       answer,
            "sources":      sources,
            "agent":        "research",
            "model":        model_name,
            "context_hits": {"pubmed": len(pubmed_hits), "biomodels": len(biomodels_hits)},
        }

    finally:
        if should_close:
            await client.aclose()
